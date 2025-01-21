from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Border, Side
from locators import Locators

class ExcelAction:
    def __init__(self, tournament_name):
        self.tournament_name = tournament_name
        print('Creating workbook...')

    def prepare_workbook_template(self):
        # Cell format parameters ---
        self.thin_border = Border(
            left = Side(border_style="thin", color="000000"),
            right = Side(border_style="thin", color="000000"),
            top = Side(border_style="thin", color="000000"),
            bottom = Side(border_style="thin", color="000000")
        )
        self.yellow_fill = PatternFill(start_color='FFFF00', end_color='FFFF00',fill_type='solid')
        # Sheet creation ---
        self.workbook = Workbook()
        for sheet in Locators.excel_sheets.keys():
            self.workbook.create_sheet(title=sheet)
        self.workbook.remove(self.workbook.active)     
        # Column heading creation ---
        for sheet_key, sheet_value in Locators.excel_sheets.items():
            current_sheet = self.workbook[sheet_key]
            for index in range(len(sheet_value)):
                current_cell = current_sheet.cell(row=1, column=index+1, value=sheet_value[index])
                current_cell.font = Font(bold=True)
                current_cell.fill = self.yellow_fill
                current_cell.border = self.thin_border
        self.workbook.save(f"{self.tournament_name}.xlsx")
        print("Workbook created and saved successfully.")
        return True

    def write_stats(self, stat_list, sheetname):
        self.stat_list = stat_list
        self.sheetname = sheetname
        current_sheet = self.workbook[sheetname]
        last_used_row = self.get_last_used_row(sheetname)
        count = 1
        for stat in stat_list:
            if sheetname in ('match_stats', 'innings_wkfall', 'player_info'):
                current_sheet.cell(row=last_used_row+1, column=count, value=stat)
            elif sheetname in ('mvp', 'overs', 'bowl_stat', 'commentary'):
                for i in range(len(stat)):
                    current_sheet.cell(row=last_used_row+i+1, column=count, value=stat[i])
            elif sheetname == 'bat_stat': 
                try: 
                    for i in range(len(stat_list[-1])):
                        current_sheet.cell(row=last_used_row+i+1, column=count, value=stat[i])
                except:
                    print('exception passing')
            count += 1
            self.workbook.save(f"{self.tournament_name}.xlsx")        

    def get_last_used_row(self, sheetname):
        self.sheetname = sheetname
        current_sheet = self.workbook[sheetname]
        # Iterate from the last row to the first row
        for row in reversed(range(1, current_sheet.max_row + 1)):
            for col in current_sheet.iter_cols(1, current_sheet.max_column):
                if col[row-1].value is not None:
                    return row
        return 0
